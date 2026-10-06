"""Convert a 神娃人物設定 pptx (Leo's layout) into a Notion page via REST.

Usage:
  python ppt2notion.py parse  <deck.pptx>                 # print parsed structure as JSON
  python ppt2notion.py build  <deck.pptx> <parent_page_id> [--title T] [--no-images]

Uses the 一部 (pm1) Notion proxy credential via notion_api / notion_rest
(see plan-dept14-writer/references/notion-access.md). Images are uploaded with the File Upload API.
"""
import io
import json
import re
import sys
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Emu

sys.path.insert(0, str(Path(__file__).resolve().parent))
from notion_api import Notion, find_token  # noqa: E402

FULLWIDTH = str.maketrans("０１２３４５６７８９", "0123456789")

ALL_PARTS = ["五官", "臉型", "服飾", "髮型", "髮飾", "眼鏡", "鬍子", "裝飾", "耳環", "武器", "特效", "翅膀"]


# ----------------------------------------------------------------- parsing
def shapes_flat(shapes):
    for sh in shapes:
        yield sh
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from shapes_flat(sh.shapes)


def inches(v):
    return round(Emu(v).inches, 1)


def table_rows(sh):
    return [[c.text.strip() for c in r.cells] for r in sh.table.rows]


def pairs(row, skip=("其他",)):
    """[name, mark, name, mark...] -> {name: mark}, starting at index 2."""
    out = {}
    cells = row[2:]
    for i in range(0, len(cells) - 1, 2):
        name, mark = cells[i], cells[i + 1]
        if name and name not in skip:
            out[name] = mark
    return out


def parse_character_slide(slide, idx):
    info = {"kind": "character", "slide": idx, "images": [], "notes": []}
    numbers = []  # (left, top, number)
    for sh in shapes_flat(slide.shapes):
        if getattr(sh, "has_table", False) and sh.has_table:
            rows = table_rows(sh)
            w, h = inches(sh.width), inches(sh.height)
            if rows and rows[0] and rows[0][0].startswith("【"):
                head = rows[0][0]
                m = re.match(r"【(.+?)\s*[–-]\s*(.+?)】(.*)", head)
                info["group"], info["gender"], info["head_note"] = (m.group(1).strip(), m.group(2).strip(), m.group(3).strip()) if m else (head, "", "")
                setting = rows[1][0] if len(rows) > 1 else ""
                m2 = re.search(r"形象設定[：:]\s*(.+?)\s{1,}配色[：:]\s*(.+)", setting)
                info["persona"] = m2.group(1).strip() if m2 else setting
                old = {}
                old.update(pairs(rows[0]))
                if len(rows) > 2:
                    old.update(pairs(rows[2]))
                new = pairs(rows[3]) if len(rows) > 3 else {}
                info["old_parts"], info["new_parts"] = old, new
            elif w < 2 and h < 1 and inches(sh.top) < 2:
                info["main_color"] = " / ".join(c for c in rows[0] if c)
            elif inches(sh.top) == 0 and w > 12:
                info["tab"] = " / ".join(c for c in rows[0] if c)
        elif sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
            info["images"].append({"left": inches(sh.left), "top": inches(sh.top), "ext": sh.image.ext, "blob": sh.image.blob})
        elif sh.has_text_frame:
            txt = sh.text_frame.text.strip()
            if not txt:
                continue
            if txt.startswith("重點需求"):
                paras = []
                for p in sh.text_frame.paragraphs:
                    for line in re.split(r"[\x0b\n]", p.text):
                        line = re.sub(r"^重點需求[：:]\s*", "", line.strip())
                        if line:
                            paras.append(line)
                info["requirements"] = paras
            elif inches(sh.width) <= 0.5 and inches(sh.height) <= 0.5:
                n = txt.translate(FULLWIDTH)
                if n.isdigit():
                    numbers.append((inches(sh.left), inches(sh.top), int(n)))
    # match numbers to images by matching top-left corner
    for im in info["images"]:
        best = None
        for (l, t, n) in numbers:
            d = abs(l - im["left"]) + abs(t - im["top"])
            if d < 0.35 and (best is None or d < best[0]):
                best = (d, n)
        im["no"] = best[1] if best else None
    info["images"].sort(key=lambda x: (x["no"] is None, x["no"] or 0, x["left"]))
    for i, im in enumerate(info["images"], 1):
        if im["no"] is None:
            im["no"] = i
    return info


def parse_scene_slide(slide, idx):
    info = {"kind": "scene", "slide": idx, "items": []}
    pics = []
    for sh in shapes_flat(slide.shapes):
        if getattr(sh, "has_table", False) and sh.has_table:
            rows = table_rows(sh)
            if inches(sh.top) == 0:
                info["tab"] = " / ".join(c for c in rows[0] if c)
            elif len(rows) >= 2:
                for col in range(len(rows[0])):
                    name = rows[0][col]
                    desc = rows[1][col] if col < len(rows[1]) else ""
                    if name:
                        info["items"].append({"name": name, "desc": desc, "images": []})
        elif sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
            pics.append({"left": inches(sh.left), "ext": sh.image.ext, "blob": sh.image.blob})
    pics.sort(key=lambda p: p["left"])
    # left half -> item 0, right half -> item 1
    for p in pics:
        tgt = 0 if (p["left"] < 6.3 or len(info["items"]) == 1) else min(1, len(info["items"]) - 1)
        if info["items"]:
            info["items"][tgt]["images"].append(p)
    return info


def parse_overview(slide):
    info = {"kind": "overview", "fields": [], "colors": {}}
    cells = []
    for sh in shapes_flat(slide.shapes):
        if getattr(sh, "has_table", False) and sh.has_table:
            rows = table_rows(sh)
            w = inches(sh.width)
            if w > 10 and inches(sh.top) < 1:
                info["fields"] = rows
            elif w > 10:
                info["grid"] = rows
            elif w < 2:
                cells.append((inches(sh.left), inches(sh.top), " / ".join(c for c in rows[0] if c)))
    # map small color cells to grid: x -> column (新手..BOSS2), y -> row (男 4.6 / 女 6.2)
    cols = ["新手", "一般 1", "一般 2", "儲值", "BOSS 1", "BOSS 2"]
    xs = [3.9, 5.5, 7.1, 8.7, 10.3, 11.9]
    for (x, y, val) in cells:
        ci = min(range(6), key=lambda i: abs(xs[i] - x))
        g = "男" if y < 5.4 else "女"
        info["colors"][f"{cols[ci]}|{g}"] = val
    return info


def parse_deck(path):
    prs = Presentation(path)
    slides = list(prs.slides)
    deck = {"title": "", "author": "", "sections": []}
    for sh in slides[0].shapes:
        if sh.has_text_frame and sh.text_frame.text.strip():
            t = sh.text_frame.text.strip()
            if "大活動" in t and not deck["title"]:
                deck["title"] = t
            elif re.match(r"^[A-Za-z.]+$", t):
                deck["author"] = t
    deck["overview"] = parse_overview(slides[1])
    for i, s in enumerate(slides[2:], 3):
        texts = " ".join(sh.text_frame.text for sh in shapes_flat(s.shapes) if sh.has_text_frame)
        tables = [table_rows(sh) for sh in shapes_flat(s.shapes) if getattr(sh, "has_table", False) and sh.has_table]
        has_char = any(r and r[0] and r[0][0].startswith("【") and "部件" in " ".join(r[0]) for r in tables)
        if has_char:
            deck["sections"].append(parse_character_slide(s, i))
        else:
            deck["sections"].append(parse_scene_slide(s, i))
    return deck


# ----------------------------------------------------------------- notion
_client: Notion | None = None


def client() -> Notion:
    """Lazily built 一部 (pm1) client: 神娃 pages live in the 一部 workspace."""
    global _client
    if _client is None:
        _client = Notion(find_token(dept="一部"))
    return _client


def api(method: str, path: str, body: object = None, raw: bytes | None = None,
        headers: dict[str, str] | None = None) -> dict:
    return client().call(method, path, body=body, raw=raw, headers=headers)


def upload_image(blob: bytes, ext: str, name: str) -> str:
    return client().upload_image(blob, ext, name)


def rt(text, bold=False, color=None, code=False):
    out = []
    for chunk in [text[i:i + 1900] for i in range(0, max(len(text), 1), 1900)]:
        ann = {}
        if bold:
            ann["bold"] = True
        if code:
            ann["code"] = True
        if color:
            ann["color"] = color
        item = {"type": "text", "text": {"content": chunk}}
        if ann:
            item["annotations"] = ann
        out.append(item)
    return out


def heading(level, text):
    return {"object": "block", "type": f"heading_{level}", f"heading_{level}": {"rich_text": rt(text)}}


def para(text, **kw):
    return {"object": "block", "type": "paragraph", "paragraph": {"rich_text": rt(text, **kw)}}


def bullet(text, **kw):
    return {"object": "block", "type": "bulleted_list_item", "bulleted_list_item": {"rich_text": rt(text, **kw)}}


def table(rows, header=True):
    width = max(len(r) for r in rows)
    children = []
    for r in rows:
        cells = []
        for c in list(r) + [""] * (width - len(r)):
            if isinstance(c, tuple):
                cells.append(rt(c[0], **c[1]))
            else:
                cells.append(rt(str(c)) if c != "" else [])
        children.append({"type": "table_row", "table_row": {"cells": cells}})
    return {"object": "block", "type": "table", "table": {"table_width": width, "has_column_header": header, "has_row_header": False, "children": children}}


def image_block(file_id, caption):
    return {"object": "block", "type": "image", "image": {"type": "file_upload", "file_upload": {"id": file_id}, "caption": rt(caption)}}


def divider():
    return {"object": "block", "type": "divider", "divider": {}}


def mark(v):
    if v in ("√", "✓", "✔", "V", "v"):
        return ("✅", {})
    if v == "":
        return ("—", {"color": "gray"})
    return (v, {"color": "orange"})


def parts_table(info):
    old, new = info.get("old_parts", {}), info.get("new_parts", {})
    rows = [[("部件", {"bold": True}), ("舊神娃部件", {"bold": True}), ("新人物部件", {"bold": True})]]
    for p in ALL_PARTS:
        o = mark(old[p]) if p in old else ("", {})
        n = mark(new[p]) if p in new else ("", {})
        if p not in old and p not in new:
            continue
        rows.append([p, o, n])
    return table(rows)


def column_list(left, right):
    col = lambda kids: {"object": "block", "type": "column", "column": {"children": kids}}
    return {"object": "block", "type": "column_list", "column_list": {"children": [col(left), col(right)]}}


def build_blocks_columns(deck, with_images=True, log=print):
    """Same content as build_blocks, but every character / scene item is a 2-column row: text left, images right."""
    ov = deck["overview"]
    blocks = [heading(1, "活動總覽")]
    for r in ov.get("fields", []):
        cells = [c for c in r if c]
        if cells:
            blocks.append(bullet(" ｜ ".join(cells)))
    if ov.get("colors"):
        cols = ["新手", "一般 1", "一般 2", "儲值", "BOSS 1", "BOSS 2"]
        rows = [[("主色", {"bold": True})] + [(c, {"bold": True}) for c in cols]]
        for g in ("男", "女"):
            rows.append([(g, {"bold": True})] + [ov["colors"].get(f"{c}|{g}", "") for c in cols])
        blocks.append(table(rows))
    blocks.append(divider())

    current_group = None
    for sec in deck["sections"]:
        if sec["kind"] == "character":
            grp = sec.get("group", "")
            if grp != current_group:
                blocks.append(heading(1, grp + ("　" + sec.get("head_note", "") if sec.get("head_note") else "")))
                current_group = grp
            left = [heading(2, f"{sec.get('gender', '')}｜{sec.get('persona', '')}"),
                    para(f"主色：{sec.get('main_color', '')}", bold=True),
                    parts_table(sec)]
            reqs = sec.get("requirements", [])
            if reqs:
                left.append(para("重點需求", bold=True))
                left += [bullet(r) for r in reqs]
            right = []
            if with_images:
                for im in sec["images"]:
                    name = f"s{sec['slide']:02d}_{grp}_{sec.get('gender', '')}_圖{im['no']}"
                    fid = upload_image(im["blob"], im["ext"], name)
                    log(f"  uploaded {name}")
                    right.append(image_block(fid, f"圖{im['no']}"))
            if not right:
                right = [para("（無參考圖）", color="gray")]
            blocks.append(column_list(left, right))
        else:
            blocks.append(heading(2, sec.get("tab", "前景／背景").split(" / ")[0]))
            for it in sec["items"]:
                left = [para(it["name"], bold=True)]
                if it["desc"]:
                    left.append(bullet(it["desc"]))
                right = []
                if with_images:
                    for k, im in enumerate(it["images"], 1):
                        fid = upload_image(im["blob"], im["ext"], f"s{sec['slide']:02d}_{it['name']}_{k}")
                        log(f"  uploaded s{sec['slide']:02d} {it['name']} #{k}")
                        right.append(image_block(fid, it["name"]))
                if not right:
                    right = [para("（無參考圖）", color="gray")]
                blocks.append(column_list(left, right))
            blocks.append(divider())
    return blocks


def build_blocks(deck, with_images=True, log=print):
    ov = deck["overview"]
    blocks = [heading(1, "活動總覽")]
    fields = []
    for r in ov.get("fields", []):
        cells = [c for c in r if c]
        if cells:
            fields.append(" ｜ ".join(cells))
    for f in fields:
        blocks.append(bullet(f))
    if ov.get("colors"):
        cols = ["新手", "一般 1", "一般 2", "儲值", "BOSS 1", "BOSS 2"]
        rows = [[("主色", {"bold": True})] + [(c, {"bold": True}) for c in cols]]
        for g in ("男", "女"):
            rows.append([(g, {"bold": True})] + [ov["colors"].get(f"{c}|{g}", "") for c in cols])
        blocks.append(table(rows))
    blocks.append(divider())

    current_group = None
    for sec in deck["sections"]:
        if sec["kind"] == "character":
            grp = sec.get("group", "")
            if grp != current_group:
                blocks.append(heading(1, grp + ("　" + sec.get("head_note", "") if sec.get("head_note") else "")))
                current_group = grp
            blocks.append(heading(2, f"{sec.get('gender', '')}｜{sec.get('persona', '')}"))
            blocks.append(para(f"主色：{sec.get('main_color', '')}", bold=True))
            blocks.append(parts_table(sec))
            reqs = sec.get("requirements", [])
            if reqs:
                blocks.append(para("重點需求", bold=True))
                for r in reqs:
                    r = re.sub(r"^重點需求[：:]\s*", "", r)
                    if r:
                        blocks.append(bullet(r))
            if with_images:
                for im in sec["images"]:
                    name = f"s{sec['slide']:02d}_{grp}_{sec.get('gender', '')}_圖{im['no']}"
                    fid = upload_image(im["blob"], im["ext"], name)
                    log(f"  uploaded {name}")
                    blocks.append(image_block(fid, f"圖{im['no']}"))
        else:
            blocks.append(heading(2, sec.get("tab", "前景／背景").split(" / ")[0]))
            for it in sec["items"]:
                blocks.append(para(it["name"], bold=True))
                if it["desc"]:
                    blocks.append(bullet(it["desc"]))
                if with_images:
                    for k, im in enumerate(it["images"], 1):
                        fid = upload_image(im["blob"], im["ext"], f"s{sec['slide']:02d}_{it['name']}_{k}")
                        log(f"  uploaded s{sec['slide']:02d} {it['name']} #{k}")
                        blocks.append(image_block(fid, it["name"]))
            blocks.append(divider())
    return blocks


def create_page(parent_id, title, blocks, log=print, deep=False):
    # Page creation allows only two nesting levels; column_list>column>table>table_row needs three,
    # which the append endpoint accepts. So for deep layouts create the page with its first block only.
    split = 1 if deep else 100
    first, rest = blocks[:split], blocks[split:]
    page = api("POST", "/pages", {"parent": {"page_id": parent_id}, "properties": {"title": {"title": rt(title)}}, "children": first})
    pid = page["id"]
    log(f"page created {page['url']}")
    while rest:
        chunk, rest = rest[:100], rest[100:]
        api("PATCH", f"/blocks/{pid}/children", {"children": chunk})
        log(f"  appended {len(chunk)} blocks")
    return page


def count_blocks(pid):
    n, cur = 0, None
    while True:
        q = f"/blocks/{pid}/children?page_size=100" + (f"&start_cursor={cur}" if cur else "")
        d = api("GET", q)
        n += len(d["results"])
        if not d.get("has_more"):
            return n
        cur = d["next_cursor"]


def strip_blobs(deck):
    d = json.loads(json.dumps(deck, default=lambda o: f"<{len(o)} bytes>" if isinstance(o, (bytes, bytearray)) else str(o), ensure_ascii=False))
    return d


if __name__ == "__main__":
    cmd = sys.argv[1]
    deck = parse_deck(sys.argv[2])
    if cmd == "parse":
        print(json.dumps(strip_blobs(deck), ensure_ascii=False, indent=1))
    elif cmd == "build":
        parent = sys.argv[3]
        title = None
        with_images = "--no-images" not in sys.argv
        if "--title" in sys.argv:
            title = sys.argv[sys.argv.index("--title") + 1]
        title = title or deck["title"]
        columns = "--columns" in sys.argv
        builder = build_blocks_columns if columns else build_blocks
        blocks = builder(deck, with_images=with_images)
        print(f"total blocks: {len(blocks)}")
        page = create_page(parent, title, blocks, deep=columns)
        print("verify children:", count_blocks(page["id"]))
        print("URL:", page["url"])
